"""Dump a .docx / .xlsx / .md requirement doc to plain text on stdout.

Usage: python doc_to_text.py <file>
Needs python-docx (docx) and openpyxl (xlsx).
"""
import sys
from pathlib import Path


def para_text(p):
    # Keep hyperlink targets: specs often hide the Figma/reference link behind text.
    parts = []
    for item in p.iter_inner_content():
        url = getattr(item, "url", None)
        parts.append(f"{item.text} ({url})" if url else item.text)
    return "".join(parts).replace("\n", " ").strip()


def docx_text(path):
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = Document(path)
    out = []
    for section in doc.sections:
        for part, label in ((section.header, "Header"), (section.footer, "Footer")):
            text = " ".join(para_text(p) for p in part.paragraphs).strip()
            if text and f"[{label}] {text}" not in out:
                out.append(f"[{label}] {text}")
    # Walk body in order so tables stay next to the paragraphs around them.
    for child in doc.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            p = Paragraph(child, doc)
            text = para_text(p)
            if not text:
                continue
            style = (p.style.name or "") if p.style is not None else ""
            if style.startswith("Heading"):
                level = style.replace("Heading", "").strip()
                out.append("#" * (int(level) if level.isdigit() else 1) + " " + text)
            elif "List" in style:
                out.append("- " + text)
            else:
                out.append(text)
        elif tag == "tbl":
            for row in Table(child, doc).rows:
                out.append("| " + " | ".join(" ".join(para_text(p) for p in c.paragraphs).strip() for c in row.cells) + " |")
            out.append("")
    comments = list(doc.comments)
    if comments:
        out.append("\n## Comments trong tài liệu")
        out.extend(f"- [{c.author}] {c.text}" for c in comments)
    if any(r for r in doc.inline_shapes):
        out.append(f"[{len(doc.inline_shapes)} hình ảnh nhúng — không đọc được nội dung ảnh]")
    return "\n".join(out)


def xlsx_text(path):
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True, read_only=True)
    out = []
    for ws in wb.worksheets:
        out.append(f"## Sheet: {ws.title}")
        for row in ws.iter_rows(values_only=True):
            cells = ["" if v is None else str(v).replace("\n", " ") for v in row]
            if any(cells):
                out.append("| " + " | ".join(cells) + " |")
        out.append("")
    return "\n".join(out)


def main():
    if len(sys.argv) != 2:
        sys.exit("Usage: python doc_to_text.py <file.docx|file.xlsx|file.md>")
    path = Path(sys.argv[1])
    ext = path.suffix.lower()
    if ext == ".docx":
        text = docx_text(path)
    elif ext in (".xlsx", ".xlsm"):
        text = xlsx_text(path)
    elif ext in (".md", ".txt"):
        text = path.read_text(encoding="utf-8")
    else:
        sys.exit(f"Unsupported file type: {ext}")
    sys.stdout.reconfigure(encoding="utf-8")
    print(text)


if __name__ == "__main__":
    main()
